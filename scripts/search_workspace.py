import docx
import pypdf
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

search_query = "reducir el riesgo"
workspace_dir = r"c:\Users\Admin\Desktop\Proyect-ISTQ\ISTQB"

print(f"Searching for '{search_query}' in {workspace_dir}...")

def search_docx(path):
    try:
        doc = docx.Document(path)
        for i, p in enumerate(doc.paragraphs):
            if search_query in p.text.lower():
                print(f"[DOCX] Found in {os.path.basename(path)} L{i}: {p.text}")
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    if search_query in cell.text.lower():
                        print(f"[DOCX-Table] Found in {os.path.basename(path)} Table{t_idx} Row{r_idx} Col{c_idx}: {cell.text[:150]}")
    except Exception as e:
        pass

def search_pdf(path):
    try:
        reader = pypdf.PdfReader(path)
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if search_query in text.lower():
                print(f"[PDF] Found in {os.path.basename(path)} Page{i+1}: ... {text[text.lower().index(search_query)-100:text.lower().index(search_query)+150]} ...")
    except Exception as e:
        pass

for root, dirs, files in os.walk(workspace_dir):
    for fname in files:
        path = os.path.join(root, fname)
        if fname.endswith(".docx"):
            search_docx(path)
        elif fname.endswith(".pdf"):
            search_pdf(path)

print("Search complete.")
